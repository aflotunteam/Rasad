"""Hisobotni PDF, XLSX va CSV ko'rinishida eksport qilish."""

from __future__ import annotations

import csv
import io
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.services.common import CONFIDENCE_LABELS, EXPERT_STATUS_LABELS, LEVEL_LABELS

FONT_DIR = Path(__file__).parent / "fonts"
NAVY = colors.HexColor("#0B1F3A")
MUTED = colors.HexColor("#607087")
BORDER = colors.HexColor("#DDE7F0")
SYNTHETIC = "NAMOYISH MA’LUMOTLARI — SINTETIK"

_fonts_ready = False


def _fonts() -> tuple[str, str]:
    global _fonts_ready
    if not _fonts_ready:
        pdfmetrics.registerFont(TTFont("DejaVu", str(FONT_DIR / "DejaVuSans.ttf")))
        pdfmetrics.registerFont(TTFont("DejaVu-Bold", str(FONT_DIR / "DejaVuSans-Bold.ttf")))
        _fonts_ready = True
    return "DejaVu", "DejaVu-Bold"


def num(v, digits: int = 1) -> str:
    if v is None:
        return "—"
    return f"{v:.{digits}f}".replace(".", ",")


def meta_rows(meta: dict) -> list[list[str]]:
    return [
        ["Hisobot raqami", meta["number"]],
        ["Hisobot turi", meta["type_label"]],
        ["Yaratilgan vaqt", meta["created_at"][:19].replace("T", " ")],
        ["Ma’lumot davri", f"{meta['period_start']} — {meta['period_end']}"],
        ["Model versiyasi", meta["model_version"]],
        ["Ma’lumot versiyasi", meta["data_version"]],
        ["Yaratgan foydalanuvchi", meta["created_by"]],
    ]


def table_rows(content: dict) -> tuple[list[str], list[list]]:
    if content["kind"] == "subject":
        header = ["Omil", "Joriy", "Kutilgan oraliq", "Og‘ish", "Ta’sir (ball)", "Manba"]
        rows = [[f["factor"], _val(f["current_value"], f["unit"]),
                 f"{_val(f['baseline_low'], f['unit'])} … {_val(f['baseline_high'], f['unit'])}",
                 LEVEL_LABELS.get(f["deviation"], f["deviation"]), num(f["impact"]), f["source"]]
                for f in content["factors"]]
        return header, rows
    header = ["Ichki kod", "Hudud", "Soha", "Xavf bahosi", "Daraja", "Ishonch", "Xavf turi", "Ma’lumot sifati",
              "Ekspert holati"]
    rows = [[t["code"], t["region"], t["sector"], num(t["score"]), t["level"], t["confidence"], t["risk_type"],
             num(t["dq_score"]), t["expert_status"]] for t in content["top"]]
    return header, rows


def _val(v, unit: str) -> str:
    if v is None:
        return "—"
    if unit == "%":
        return f"{num(v)}%"
    if unit == "ta":
        return f"{int(v)} ta"
    return num(v, 3)


def to_pdf(meta: dict, content: dict) -> bytes:
    regular, bold = _fonts()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm, topMargin=16 * mm,
                            bottomMargin=16 * mm, title=meta["title"], author="RASAD")
    h1 = ParagraphStyle("h1", fontName=bold, fontSize=15, leading=19, textColor=NAVY, spaceAfter=4)
    h2 = ParagraphStyle("h2", fontName=bold, fontSize=11, leading=14, textColor=NAVY, spaceBefore=10, spaceAfter=4)
    body = ParagraphStyle("body", fontName=regular, fontSize=9, leading=13, alignment=TA_LEFT)
    small = ParagraphStyle("small", fontName=regular, fontSize=7.5, leading=10, textColor=MUTED)
    badge = ParagraphStyle("badge", fontName=bold, fontSize=8, leading=10, textColor=colors.HexColor("#216DF3"))

    story = [Paragraph("RASAD", badge), Paragraph(meta["title"], h1), Paragraph(SYNTHETIC, badge), Spacer(1, 4)]
    mt = Table(meta_rows(meta), colWidths=[45 * mm, 120 * mm])
    mt.setStyle(TableStyle([
        ("FONT", (0, 0), (-1, -1), regular, 8.5), ("TEXTCOLOR", (0, 0), (0, -1), MUTED),
        ("LINEBELOW", (0, 0), (-1, -1), 0.3, BORDER), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    story += [mt]

    if content["kind"] == "subject":
        r = content["risk"]
        s = content["subject"]
        story.append(Paragraph("Umumiy ma’lumot", h2))
        info = [["Ichki kod", s["code"]], ["Hudud", s["region"]["name"]], ["Faoliyat turi", s["sector"]["name"]],
                ["STIR", s["stir"]], ["Xavf bahosi", f"{num(r['score'])} / 100 ({LEVEL_LABELS[r['level']]})"],
                ["Ishonch darajasi", CONFIDENCE_LABELS[r["confidence"]]],
                ["Ma’lumot sifati", f"{num(r['dq_score'])} / 100"],
                ["Asosiy xavf turi", r["primary_risk_type"]],
                ["Ekspert holati", EXPERT_STATUS_LABELS.get(r["expert_status"], r["expert_status"])]]
        it = Table(info, colWidths=[45 * mm, 120 * mm])
        it.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), regular, 9), ("TEXTCOLOR", (0, 0), (0, -1), MUTED),
                                ("LINEBELOW", (0, 0), (-1, -1), 0.3, BORDER)]))
        story.append(it)
        story.append(Paragraph("Xavf sabablari", h2))
    else:
        sm = content["summary"]
        story.append(Paragraph("Umumiy ko‘rsatkichlar", h2))
        it = Table([["Tahlil qilingan subyektlar", str(sm["subjects"])],
                    ["Yuqori ustuvorlikdagi holatlar", str(sm["high"])],
                    ["O‘rta ustuvorlikdagi holatlar", str(sm["medium"])],
                    ["O‘rtacha ma’lumot sifati", f"{num(sm['avg_dq'])} / 100"],
                    ["Xavf chegaralari (MVP, sozlanadigan)", f"{int(sm['thresholds']['low'])} / {int(sm['thresholds']['high'])}"]],
                   colWidths=[70 * mm, 95 * mm])
        it.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), regular, 9), ("TEXTCOLOR", (0, 0), (0, -1), MUTED),
                                ("LINEBELOW", (0, 0), (-1, -1), 0.3, BORDER)]))
        story.append(it)
        story.append(Paragraph("Eng yuqori ustuvorlikdagi holatlar", h2))

    header, rows = table_rows(content)
    wrapped = [[Paragraph(str(c), small) for c in header]] + [[Paragraph(str(c), small) for c in row] for row in rows]
    t = Table(wrapped, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F5F9FD")), ("GRID", (0, 0), (-1, -1), 0.3, BORDER),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story.append(t if rows else Paragraph("Ma’lumot mavjud emas", body))

    if content["kind"] == "subject":
        exp = content["explanation"]
        story.append(Paragraph("Sun’iy intellekt izohi", h2))
        for ptxt in exp["paragraphs"]:
            story.append(Paragraph(ptxt, body))
            story.append(Spacer(1, 3))
        story.append(Paragraph(exp["basis_note"], small))
        story.append(Paragraph("Ekspert qarorlari", h2))
        if content["decisions"]:
            for d in content["decisions"]:
                story.append(Paragraph(
                    f"{d['created_at'][:16].replace('T', ' ')} · {d['user']['full_name']} · "
                    f"{EXPERT_STATUS_LABELS.get(d['decision'], d['decision'])}: {d['comment']}", body))
        else:
            story.append(Paragraph("Ekspert qarori hali qabul qilinmagan.", body))

    story += [Spacer(1, 10), Paragraph(meta["disclaimer"], small), Paragraph(SYNTHETIC, small)]
    doc.build(story)
    return buf.getvalue()


def to_xlsx(meta: dict, content: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Hisobot"
    bold = Font(bold=True, color="0B1F3A")
    ws.append(["RASAD — " + meta["title"]])
    ws["A1"].font = Font(bold=True, size=14, color="0B1F3A")
    ws.append([SYNTHETIC])
    ws.append([])
    for row in meta_rows(meta):
        ws.append(row)
        ws.cell(ws.max_row, 1).font = Font(color="607087")
    ws.append([])
    header, rows = table_rows(content)
    ws.append(header)
    for c in ws[ws.max_row]:
        c.font = bold
        c.fill = PatternFill("solid", fgColor="F5F9FD")
    for r in rows:
        ws.append(r)
    ws.append([])
    ws.append([meta["disclaimer"]])
    for col, width in zip("ABCDEFGHI", [34, 16, 22, 14, 12, 26, 12, 14, 24]):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_csv(meta: dict, content: dict) -> bytes:
    buf = io.StringIO()
    buf.write("﻿")
    w = csv.writer(buf, delimiter=";")
    for row in meta_rows(meta):
        w.writerow(row)
    w.writerow([])
    header, rows = table_rows(content)
    w.writerow(header)
    w.writerows(rows)
    w.writerow([])
    w.writerow([meta["disclaimer"]])
    w.writerow([SYNTHETIC])
    return buf.getvalue().encode("utf-8")
